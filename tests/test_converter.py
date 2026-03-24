import os
import sqlite3
import pytest
from table2db import TableConverter
from table2db.errors import FileReadError, UnsupportedFormatError, NoDataError


@pytest.fixture
def converter():
    return TableConverter()


class TestEndToEnd:
    def test_simple(self, converter, fixture_path):
        with converter.convert(fixture_path("simple.xlsx")) as result:
            assert len(result.tables) == 1
            conn = sqlite3.connect(result.db_path)
            rows = conn.execute(f'SELECT * FROM "{result.tables[0].name}"').fetchall()
            assert len(rows) == 20
            conn.close()

    def test_merged_cells(self, converter, fixture_path):
        with converter.convert(fixture_path("merged_cells.xlsx")) as result:
            assert len(result.tables) == 1
            conn = sqlite3.connect(result.db_path)
            tbl = result.tables[0].name
            rows = conn.execute(f'SELECT * FROM "{tbl}"').fetchall()
            assert len(rows) == 12
            # Group A rows (id 1-3): label "Group A" filled, group_total 600 filled
            group_a = conn.execute(f'SELECT * FROM "{tbl}" WHERE name = ?', ("Group A",)).fetchall()
            assert len(group_a) == 3
            assert all(r[3] == 600 for r in group_a)  # group_total = 600 for all
            # Group D rows (id 10-12): standalone, each has own total
            group_d = conn.execute(f'SELECT * FROM "{tbl}" WHERE name = ?', ("Group D",)).fetchall()
            assert len(group_d) == 3
            conn.close()

    def test_multi_header(self, converter, fixture_path):
        with converter.convert(fixture_path("multi_header.xlsx")) as result:
            assert len(result.tables) == 1
            t = result.tables[0]
            # Should have compound column names from multi-level headers
            col_names = [c["name"] for c in t.columns]
            # At least one compound name should exist
            assert any("_" in name for name in col_names if name != "column_1")
            conn = sqlite3.connect(result.db_path)
            conn.close()

    def test_subtotals_filtered(self, converter, fixture_path):
        with converter.convert(fixture_path("subtotals.xlsx")) as result:
            conn = sqlite3.connect(result.db_path)
            rows = conn.execute(f'SELECT * FROM "{result.tables[0].name}"').fetchall()
            # Original: 12 data rows + 5 subtotal rows. After filtering: 12 data rows
            assert len(rows) == 12
            # Verify no subtotal keywords in data
            for row in rows:
                for cell in row:
                    if isinstance(cell, str):
                        assert cell.lower() not in ("小计", "subtotal", "合计")
            conn.close()

    def test_mixed_types(self, converter, fixture_path):
        with converter.convert(fixture_path("mixed_types.xlsx")) as result:
            conn = sqlite3.connect(result.db_path)
            rows = conn.execute(f'SELECT * FROM "{result.tables[0].name}"').fetchall()
            assert len(rows) >= 20
            conn.close()

    def test_offset_table(self, converter, fixture_path):
        with converter.convert(fixture_path("offset_table.xlsx")) as result:
            assert len(result.tables) == 1
            conn = sqlite3.connect(result.db_path)
            rows = conn.execute(f'SELECT * FROM "{result.tables[0].name}"').fetchall()
            assert len(rows) == 20
            conn.close()

    def test_multi_sheet_fk(self, converter, fixture_path):
        with converter.convert(fixture_path("multi_sheet_fk.xlsx")) as result:
            assert len(result.tables) == 3
            table_names = [t.name for t in result.tables]
            # All three sheets should become tables
            assert len(table_names) == 3

    def test_number_as_text(self, converter, fixture_path):
        with converter.convert(fixture_path("number_as_text.xlsx")) as result:
            conn = sqlite3.connect(result.db_path)
            # The "code" and "amount" columns should be detected as numeric
            t = result.tables[0]
            type_map = {c["name"]: c["type"] for c in t.columns}
            # At least code or amount should be numeric (INTEGER or REAL)
            numeric_types = {"INTEGER", "REAL"}
            has_numeric = any(v in numeric_types for k, v in type_map.items() if k != "id")
            assert has_numeric
            conn.close()

    def test_dates(self, converter, fixture_path):
        with converter.convert(fixture_path("dates_mixed.xlsx")) as result:
            t = result.tables[0]
            type_map = {c["name"]: c["type"] for c in t.columns}
            conn = sqlite3.connect(result.db_path)
            conn.close()

    def test_empty_gaps(self, converter, fixture_path):
        with converter.convert(fixture_path("empty_gaps.xlsx")) as result:
            conn = sqlite3.connect(result.db_path)
            rows = conn.execute(f'SELECT * FROM "{result.tables[0].name}"').fetchall()
            # Empty rows should be removed, leaving 20 data rows
            assert len(rows) == 20
            conn.close()

    def test_context_manager_cleanup(self, converter, fixture_path):
        with converter.convert(fixture_path("simple.xlsx")) as result:
            path = result.db_path
            assert os.path.exists(path)
        assert not os.path.exists(path)

    def test_duplicate_sheet_names(self, converter, fixture_path):
        with converter.convert(fixture_path("duplicate_sheet_names.xlsx")) as result:
            assert len(result.tables) == 2
            names = [t.name for t in result.tables]
            assert names[0] != names[1]

    def test_real_world_dirty(self, converter, fixture_path):
        """Comprehensive test with offset + merged headers + subtotals + mixed types."""
        with converter.convert(fixture_path("real_world_dirty.xlsx")) as result:
            assert len(result.tables) >= 1
            conn = sqlite3.connect(result.db_path)
            rows = conn.execute(f'SELECT * FROM "{result.tables[0].name}"').fetchall()
            # Should have some data rows (original data minus subtotals)
            assert len(rows) >= 2
            conn.close()

    def test_multi_table_sheet(self, converter, fixture_path):
        """One sheet with two tables → two DB tables."""
        with converter.convert(fixture_path("multi_table_sheet.xlsx")) as result:
            assert len(result.tables) >= 2
            conn = sqlite3.connect(result.db_path)
            try:
                for t in result.tables:
                    rows = conn.execute(f'SELECT COUNT(*) FROM "{t.name}"').fetchone()[0]
                    assert rows >= 5
            finally:
                conn.close()

    def test_csv_convert(self, converter, fixture_path):
        with converter.convert(fixture_path("students_csv.csv")) as result:
            assert len(result.tables) == 1
            conn = sqlite3.connect(result.db_path)
            try:
                rows = conn.execute(f'SELECT COUNT(*) FROM "{result.tables[0].name}"').fetchall()
                assert rows[0][0] == 20
            finally:
                conn.close()

    def test_tsv_convert(self, converter, fixture_path):
        with converter.convert(fixture_path("products_tsv.tsv")) as result:
            assert len(result.tables) == 1
            conn = sqlite3.connect(result.db_path)
            try:
                rows = conn.execute(f'SELECT COUNT(*) FROM "{result.tables[0].name}"').fetchall()
                assert rows[0][0] == 15
            finally:
                conn.close()

    def test_bytesio_convert(self, converter, fixture_path):
        """Convert from BytesIO stream instead of file path."""
        import io
        path = fixture_path("simple.xlsx")
        with open(path, "rb") as f:
            stream = io.BytesIO(f.read())
        with converter.convert(stream, file_name="simple.xlsx") as result:
            assert len(result.tables) == 1
            assert result.tables[0].row_count == 20

    def test_confidence_in_table_info(self, converter, fixture_path):
        """TableInfo should carry confidence from island detection."""
        with converter.convert(fixture_path("simple.xlsx")) as result:
            for t in result.tables:
                assert hasattr(t, "confidence")
                assert t.confidence > 0.0

    def test_multi_table_confidence(self, converter, fixture_path):
        """Multi-table sheet: each sub-table has its own confidence."""
        with converter.convert(fixture_path("multi_table_sheet.xlsx")) as result:
            assert len(result.tables) >= 2
            for t in result.tables:
                assert t.confidence > 0.0

    def test_color_mode_parameter(self, converter, fixture_path):
        """Verify color_mode parameter is accepted."""
        c = TableConverter(color_mode="value")
        with c.convert(fixture_path("simple.xlsx")) as result:
            assert len(result.tables) >= 1


class TestErrors:
    def test_file_not_found(self, converter):
        with pytest.raises(FileReadError):
            converter.convert("nonexistent.xlsx")

    def test_unsupported_format(self, converter, tmp_path):
        bad_file = tmp_path / "test.txt"
        bad_file.write_text("not excel")
        with pytest.raises(UnsupportedFormatError):
            converter.convert(str(bad_file))

    def test_empty_after_clean(self, converter, fixture_path):
        """Few data rows + many subtotal rows → data survives, totals filtered."""
        with converter.convert(fixture_path("empty_after_clean.xlsx")) as result:
            conn = sqlite3.connect(result.db_path)
            rows = conn.execute(f'SELECT * FROM "{result.tables[0].name}"').fetchall()
            assert len(rows) == 3  # 3 data rows, 4 total rows filtered
            # Verify no subtotal keywords remain
            for row in rows:
                for cell in row:
                    if isinstance(cell, str):
                        assert cell not in ("小计", "合计")
            conn.close()

    def test_no_data_error(self, converter, tmp_path):
        """Empty Excel with only headers → NoDataError."""
        from openpyxl import Workbook
        wb = Workbook()
        ws = wb.active
        ws.append(["col_a", "col_b", "col_c"])
        path = str(tmp_path / "empty.xlsx")
        wb.save(path)
        with pytest.raises(NoDataError):
            converter.convert(path)
